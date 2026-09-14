import pymupdf
import pytest
from docx import Document

from catalina_ingestion.domain.exceptions import ExtractionError
from catalina_ingestion.extractors.factory import ExtractorFactory


def test_extracts_text_document(tmp_path):
    path = tmp_path / "reglamento.txt"
    path.write_text(
        "REGLAMENTO ESTUDIANTIL\n\nLos estudiantes tienen derechos y deberes.",
        encoding="utf-8",
    )

    result = ExtractorFactory().extract(path)

    assert result.page_count == 1
    assert len(result.blocks) == 2
    assert result.blocks[0].block_type == "heading"
    assert result.blocks[1].section == "REGLAMENTO ESTUDIANTIL"


def test_splits_embedded_text_headings(tmp_path):
    path = tmp_path / "registro.txt"

    path.write_text(
        (
            "CERTIFICADOS Y CONSTANCIAS EN LÍNEA\n\n"
            "Los estudiantes pueden solicitar "
            "certificados académicos.\n\n"
            "La entrega puede tardar hasta "
            "ocho días hábiles.\n\n"
            "CUENTA INSTITUCIONAL\n"
            "Para utilizar el procedimiento "
            "se requiere una cuenta activa.\n\n"
            "VERIFICACIONES ACADÉMICAS\n"
            "Las verificaciones deben ser "
            "solicitadas formalmente."
        ),
        encoding="utf-8",
    )

    result = ExtractorFactory().extract(path)

    account_heading = next(block for block in result.blocks if block.text == "CUENTA INSTITUCIONAL")

    account_paragraph = next(
        block for block in result.blocks if block.text.startswith("Para utilizar el procedimiento")
    )

    verification_heading = next(
        block for block in result.blocks if block.text == "VERIFICACIONES ACADÉMICAS"
    )

    verification_paragraph = next(
        block for block in result.blocks if block.text.startswith("Las verificaciones deben")
    )

    assert account_heading.block_type == "heading"

    assert account_paragraph.section == "CUENTA INSTITUCIONAL"

    assert verification_heading.block_type == "heading"

    assert verification_paragraph.section == "VERIFICACIONES ACADÉMICAS"


def test_extracts_html_document(tmp_path):
    path = tmp_path / "facultad.html"
    path.write_text(
        """
        <html lang="es">
          <head><title>Facultad de Ingeniería</title></head>
          <body>
            <h1>Programas</h1>
            <p>La facultad ofrece diferentes programas.</p>
          </body>
        </html>
        """,
        encoding="utf-8",
    )

    result = ExtractorFactory().extract(path)

    assert result.title == "Facultad de Ingeniería"
    assert len(result.blocks) == 2
    assert result.blocks[0].block_type == "heading"
    assert result.blocks[1].section == "Programas"


def test_extracts_docx_document(tmp_path):
    path = tmp_path / "manual.docx"
    document = Document()
    document.add_heading("Manual del estudiante", level=1)
    document.add_paragraph("Este manual contiene información institucional.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Sede"
    table.cell(0, 1).text = "Dirección"
    table.cell(1, 0).text = "Bogotá"
    table.cell(1, 1).text = "Carrera 13"
    document.save(path)

    result = ExtractorFactory().extract(path)

    assert len(result.blocks) == 3
    assert result.blocks[0].block_type == "heading"
    assert result.blocks[2].block_type == "table"


def test_extracts_pdf_with_page_number(tmp_path):
    path = tmp_path / "reglamento.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Reglamento institucional",
        fontsize=18,
    )
    page.insert_text(
        (72, 120),
        "Este documento contiene las normas de la universidad.",
        fontsize=11,
    )
    document.save(path)
    document.close()

    result = ExtractorFactory().extract(path)

    assert result.page_count == 1
    assert result.blocks
    assert all(block.page_number == 1 for block in result.blocks)


def test_detects_pdf_that_requires_ocr(tmp_path):
    path = tmp_path / "escaneado.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(path)
    document.close()

    with pytest.raises(ExtractionError) as captured:
        ExtractorFactory().extract(path)

    assert captured.value.code == "PDF_REQUIERE_OCR"


def test_rejects_unknown_extension(tmp_path):
    path = tmp_path / "datos.xyz"
    path.write_text("contenido", encoding="utf-8")

    with pytest.raises(ExtractionError) as captured:
        ExtractorFactory().extract(path)

    assert captured.value.code == "EXTRACTOR_NOT_FOUND"
